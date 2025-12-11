import streamlit as st
from docx import Document
from datetime import datetime
import io
import os # Pastikan os ada di sini

# --- Fungsi untuk memproses dokumen ---
def generate_docx(data):
    # Dapatkan path absolut template
    script_dir = os.path.dirname(os.path.abspath(__file__))
    template_path = os.path.join(script_dir, 'template_surat.docx')
    
    # Debugging: Tampilkan path di konsol/log Streamlit untuk verifikasi
    print(f"Mencari template di: {template_path}") 

    if not os.path.exists(template_path):
        st.error(f"Template tidak ditemukan. Mencari di: {template_path}")
        return None

    # Muat template
    try:
        document = Document(template_path)
    except Exception as e:
        # Jika loading gagal, tampilkan error yang jelas
        st.error(f"Gagal memuat dokumen Word. Cek apakah file template valid (.docx): {e}")
        return None 
    
    # ... (Sisa kode untuk placeholder, tidak perlu diubah) ...
    tanggal_proses = datetime.now().strftime("%d-%m-%Y")

    placeholders = {
        '<<NAMA_LENGKAP>>': data['nama_lengkap'],
        '<<ALAMAT_SURAT>>': data['alamat_surat'],
        '<<TANGGAL_PROSES>>': tanggal_proses
    }

    for paragraph in document.paragraphs:
        for key, value in placeholders.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)
                
    # Simpan dan kembalikan file_stream
    file_stream = io.BytesIO()
    document.save(file_stream)
    file_stream.seek(0)
    
    return file_stream
# --- Tampilan Streamlit ---
st.title("📄 Pembuat Dokumen Otomatis")
st.markdown("Isi formulir di bawah ini untuk membuat dokumen Word (.docx) secara instan.")

# Membuat formulir input
with st.form("document_form"):
    nama_lengkap = st.text_input("Nama Lengkap", placeholder="Masukkan nama lengkap Anda")
    alamat_surat = st.text_area("Alamat Surat", placeholder="Masukkan alamat lengkap")
    
    # Tombol submit formulir
    submitted = st.form_submit_button("Generate Dokumen")

if submitted:
    if nama_lengkap and alamat_surat:
        
        # 1. Kumpulkan data
        input_data = {
            'nama_lengkap': nama_lengkap,
            'alamat_surat': alamat_surat
        }
        
        # 2. Hasilkan dokumen
        docx_file = generate_docx(input_data)
        
        if docx_file:
            # 3. Buat nama file yang rapi
            safe_name = nama_lengkap.replace(" ", "_")
            file_name = f"Dokumen_{safe_name}_{datetime.now().strftime('%Y%m%d')}.docx"
            
            st.success("Dokumen berhasil dibuat!")
            
            # 4. Tampilkan tombol unduh
            st.download_button(
                label="📥 Klik untuk Unduh Dokumen",
                data=docx_file,
                file_name=file_name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
            st.info(f"Dokumen bernama `{file_name}` telah siap.")
            
    else:
        st.warning("Mohon isi semua bidang yang diperlukan.")

# Tambahkan petunjuk tambahan
st.markdown("---")
st.subheader("💡 Cara Kerja:")
st.text("Aplikasi ini menggunakan Python Streamlit untuk antarmuka dan python-docx untuk mengisi 'template_surat.docx'.")